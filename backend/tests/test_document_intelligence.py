"""Document structure and complete version isolation through real analysis services."""

from datetime import datetime, timezone
import hashlib
import io
from uuid import uuid4

import pytest
from fastapi import HTTPException
from sqlalchemy import select
from docx import Document as WordDocument

from app.core.dependencies import get_current_active_user
from app.main import app
from app.modules.auth.models import User
from app.modules.organizations.models import Membership, Organization, OrganizationRole
from app.modules.documents.models import Document, DocumentVersion
from app.modules.documents.service import DocumentService
from app.modules.documents.target import AnalysisTarget
from app.modules.processing.extraction import extract_document
from app.modules.processing.models import ProcessedDocument
from app.modules.processing.service import DocumentProcessingService
from app.modules.processing.structure import build_structure, normalize_source
from app.modules.jobs.models import Job, JobType, JobStatus
from app.modules.jobs.schemas import JobCreate
from app.modules.jobs.service import JobService
from app.modules.detection.service import DetectionService
from app.modules.citations.service import CitationService
from app.modules.authorship.service import AuthorshipService
from app.modules.similarity.service import SimilarityService
from app.modules.evidence.service import EvidenceService
from app.modules.evidence.models import EvidenceNodeType, EvidenceEdgeType
from app.workers.tasks import validate_job_target


def pdf_fixture():
    """Two actual PDF pages, built without another document-generation dependency."""
    streams = [
        b"BT /F1 12 Tf 50 740 Td (First physical page. A measured result.) Tj ET",
        b"BT /F1 12 Tf 50 740 Td (Second physical page. References) Tj 0 -25 Td (Smith 2020. Source title.) Tj ET",
    ]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(streams[0])).encode()
        + b" >>\nstream\n"
        + streams[0]
        + b"\nendstream",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents 6 0 R >>",
        b"<< /Length "
        + str(len(streams[1])).encode()
        + b" >>\nstream\n"
        + streams[1]
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    result = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(result)
    result += b"xref\n0 8\n0000000000 65535 f \n"
    result += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    return (
        result
        + b"trailer\n<< /Size 8 /Root 1 0 R >>\nstartxref\n"
        + str(xref).encode()
        + b"\n%%EOF"
    )


def structure(content, extension="txt", version="v1"):
    extracted = extract_document(content, extension)
    text, document = build_structure(
        extracted, "org", "document", version, hashlib.sha256(content).hexdigest()
    )
    for collection in (
        "pages",
        "headings",
        "sections",
        "paragraphs",
        "sentences",
        "references",
        "citations",
        "tables",
    ):
        for item in document[collection]:
            assert text[item["start"] : item["end"]] == item["text"]
            source = extracted.text[item["source_start"] : item["source_end"]]
            assert normalize_source(source)[0] == item["text"]
    return text, document


def test_paragraph_sentence_reference_citation_and_unicode_source_boundaries():
    raw = (
        "Introduction\r\n\r\n🙂 Dr. Smith measured 3.5 units. The result supports this claim [1].\r\n\r\n"
        "A second paragraph cites earlier observations (Jones, 2021).\r\n\r\nReferences\r\n"
        "[1] Smith, J. (2020). Laboratory records.\r\nJones, A. (2021). Field observations."
    ).encode()
    text, doc = structure(raw)
    assert "\r" not in text and "\n\n" in text
    assert doc["page_mapping"] == "UNAVAILABLE"
    assert len(doc["references"]) == 2
    assert [c["text"] for c in doc["citations"]] == ["[1]", "(Jones, 2021)"]
    assert doc["sentences"][0]["text"] == "🙂 Dr. Smith measured 3.5 units."
    assert doc["sentences"][1]["text"] == "The result supports this claim [1]."
    assert all(c["sentence_id"] for c in doc["citations"])
    assert all(p["text"].strip() == p["text"] for p in doc["paragraphs"])
    assert len({p["id"] for p in doc["paragraphs"]}) == len(doc["paragraphs"])
    assert structure(raw)[1] == doc
    other = structure(raw, version="v2")[1]
    assert other["normalized_content_hash"] == doc["normalized_content_hash"]
    assert other["id"] != doc["id"] and other["fingerprint"] != doc["fingerprint"]
    assert not {p["id"] for p in doc["paragraphs"]} & {
        p["id"] for p in other["paragraphs"]
    }


def test_docx_preserves_headings_paragraphs_and_table_order():
    document = WordDocument()
    document.add_heading("Methods", level=1)
    document.add_paragraph("Before the table. One sentence follows.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Sample", "Value"), ("Water", "3.5")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("After the table, the observations continue.")
    buffer = io.BytesIO()
    document.save(buffer)
    text, doc = structure(buffer.getvalue(), "docx")
    assert (
        text.index("Before")
        < text.index("Sample")
        < text.index("Water")
        < text.index("After")
    )
    assert doc["headings"][0]["text"] == "Methods"
    assert doc["headings"][0]["recognition"] == "STYLE"
    assert doc["page_mapping"] == "UNAVAILABLE"
    assert len(doc["tables"]) == 1 and len(doc["tables"][0]["cells"]) == 4
    for cell in doc["tables"][0]["cells"]:
        assert text[cell["start"] : cell["end"]] == cell["text"]
        assert cell["source_path"].startswith("docx:body:")


def test_pdf_physical_pages_and_source_locations_are_exact():
    text, doc = structure(pdf_fixture(), "pdf")
    assert doc["page_mapping"] == "PHYSICAL"
    assert len(doc["pages"]) == 2
    assert doc["pages"][0]["text"].startswith("First physical page")
    assert doc["pages"][1]["text"].startswith("Second physical page")
    assert doc["pages"][0]["end"] < doc["pages"][1]["start"]
    assert any(s["page_numbers"] == [2] for s in doc["sentences"])
    assert doc["source_locations"]
    for location in doc["source_locations"]:
        assert text[location["start"] : location["end"]]
        assert len(location["bbox"]) == 4 and location["page"] in (1, 2)
    _, explicit = structure(b"First logical page.\fSecond logical page.")
    assert explicit["page_mapping"] == "EXPLICIT_BREAKS_ONLY"
    assert len(explicit["pages"]) == 2


def test_html_table_cells_and_nested_markup_keep_exact_dom_order():
    content = b"<!DOCTYPE html><html><body><!-- hidden comment --><h1>Methods</h1><p>Before <em>table</em>.</p><table><tr><th>Sample</th><th>Value</th></tr><tr><td><p>Water</p></td><td>3.5</td></tr></table><p>After table.</p></body></html>"
    text, doc = structure(content, "html")
    assert text.count("Water") == 1 and "hidden comment" not in text
    assert text.index("Before") < text.index("Water") < text.index("After")
    assert len(doc["tables"]) == 1
    assert [c["text"] for c in doc["tables"][0]["cells"]] == [
        "Sample",
        "Value",
        "Water",
        "3.5",
    ]
    assert all(
        text[c["start"] : c["end"]] == c["text"] for c in doc["tables"][0]["cells"]
    )
    assert doc["page_mapping"] == "UNAVAILABLE"


def test_authorship_consumes_frozen_structure_boundaries():
    from app.modules.authorship.intelligence import AuthorshipFeatureAnalyzer

    text, doc = structure(
        b"# Introduction\n\nDr. Smith measured 3.5 units. Another sentence follows.\n\n# References\nSmith, J. (2020). Measurements."
    )
    analyzer = AuthorshipFeatureAnalyzer()
    features = analyzer.extract(text, "version-one", doc)
    assert features.sentence_count == len(doc["sentences"])
    assert features.paragraph_count == len(doc["paragraphs"])
    assert len(doc["references"]) == 1
    baseline = analyzer.build_baseline([("version-one", text)], {"version-one": doc})
    assert (
        baseline["scalar_stats"]["mean_sentence_length"]["mean"]
        == features.scalars["mean_sentence_length"]
    )


async def identity(db):
    org = Organization(name="Structure tests", slug=str(uuid4()))
    user = User(email=f"{uuid4()}@example.invalid", hashed_password="fixture-only")
    db.add_all([org, user])
    await db.flush()
    db.add(
        Membership(
            user_id=str(user.id),
            organization_id=str(org.id),
            role=OrganizationRole.OWNER,
            is_active=True,
            joined_at=datetime.now(timezone.utc),
        )
    )
    await db.flush()
    user.current_organization_id = str(org.id)
    return org, user


async def version_input(db, org, user, text, parent=None):
    content_hash = hashlib.sha256(text.encode()).hexdigest()
    storage_path = f"tests/{uuid4()}.txt"
    if parent is not None:
        version = await DocumentService(db).create_version(
            str(parent.id),
            str(org.id),
            str(user.id),
            storage_path,
            content_hash,
            "Second revision",
        )
        return parent, version
    document = Document(
        organization_id=str(org.id),
        owner_id=str(user.id),
        filename="input.txt",
        original_filename="input.txt",
        extension="txt",
        mime_type="text/plain",
        file_size=len(text.encode()),
        storage_path=storage_path,
        sha256_fingerprint=content_hash,
    )
    db.add(document)
    await db.flush()
    version = DocumentVersion(
        document_id=str(document.id),
        version_number=1,
        storage_path=storage_path,
        content_hash=content_hash,
        sha256_fingerprint=content_hash,
        created_by_id=str(user.id),
        uploaded_by_id=str(user.id),
    )
    db.add(version)
    await db.flush()
    return document, version


async def analyze(db, org, user, document, version, text):
    parsed = await DocumentProcessingService(db).process_document(
        document, text.encode(), str(version.id)
    )
    job = await JobService(db).create_job(
        JobCreate(
            job_type=JobType.DOCUMENT_PROCESSING,
            document_id=str(document.id),
            document_version_id=str(version.id),
            input_data={
                "document_version_id": str(version.id),
                "storage_key": version.storage_path,
            },
        ),
        str(org.id),
        str(user.id),
    )
    await DetectionService(db).analyze_document(str(document.id), parsed, str(job.id))
    await CitationService(db).analyze_citations(str(document.id), parsed)
    await AuthorshipService(db).analyze_authorship(str(document.id), parsed)
    await SimilarityService(db).analyze_similarity(str(document.id), parsed)
    await EvidenceService(db).materialize_document_graph(
        str(document.id), str(org.id), str(version.id)
    )
    job.status = JobStatus.COMPLETED
    await db.flush()
    return parsed, job


async def test_v1_v2_all_analysis_api_isolation_and_historical_parser_stability(
    db_session, client, monkeypatch
):
    org, user = await identity(db_session)
    app.dependency_overrides[get_current_active_user] = lambda: user
    text1 = "🙂 River ecosystems support diverse habitats and require careful seasonal observations [1].\n\nReferences\n[1] Smith, J. (2020). Seasonal habitats."
    text2 = "Version two measures different sediment chemistry in laboratory samples (Jones, 2022).\n\nReferences\nJones, A. (2022). Mineral measurements."
    source_text = text1 + "\n\n" + text2 + "\nWorkspace comparison source."
    source, source_version = await version_input(db_session, org, user, source_text)
    source_parsed = await DocumentProcessingService(db_session).process_document(
        source, source_text.encode(), str(source_version.id)
    )
    await SimilarityService(db_session).analyze_similarity(
        str(source.id), source_parsed
    )
    document, v1 = await version_input(db_session, org, user, text1)
    p1, job1 = await analyze(db_session, org, user, document, v1, text1)
    await db_session.commit()

    async def bundle(version):
        paths = [
            f"/documents/{document.id}/content",
            f"/documents/{document.id}/structure",
            f"/detection/documents/{document.id}",
            f"/citations/documents/{document.id}/analysis",
            f"/authorship/documents/{document.id}/analysis",
            f"/similarity/documents/{document.id}/analysis",
            f"/evidence/documents/{document.id}/graph",
            f"/reports/documents/{document.id}",
        ]
        result = []
        for path in paths:
            response = await client.get(
                "/api/v1" + path, params={"document_version_id": str(version.id)}
            )
            assert response.status_code == 200, response.text
            data = response.json()
            assert data["document_version_id"] == str(version.id)
            data.pop("generated_at", None)
            result.append(data)
        graph = result[-2]
        assert graph["complete"]
        assert all(node["document_version_id"] == version.id for node in graph["nodes"])
        assert all(edge["document_version_id"] == version.id for edge in graph["edges"])
        assert all(
            segment["document_version_id"] == version.id
            for segment in result[2]["segments"]
        )
        for collection in ("findings", "citations", "references", "sources"):
            assert all(
                item["document_version_id"] == version.id
                for item in result[3][collection]
            )
        for node in graph["nodes"]:
            if node["span_text"] is not None:
                assert (
                    result[0]["content"][node["span_start"] : node["span_end"]]
                    == node["span_text"]
                )
        return result

    initial = await bundle(v1)
    assert initial[5]["matches"]["total"] > 0
    _, v2 = await version_input(db_session, org, user, text2, parent=document)
    await db_session.commit()
    # A newer unprocessed version cannot fall back to V1's text or findings.
    assert (
        await client.get(f"/api/v1/documents/{document.id}/content")
    ).status_code == 409
    assert (
        await client.get(f"/api/v1/detection/documents/{document.id}")
    ).status_code == 404
    pending = (
        await client.get(f"/api/v1/citations/documents/{document.id}/analysis")
    ).json()
    assert pending["document_version_id"] == v2.id and pending["findings"] == []
    p2, job2 = await analyze(db_session, org, user, document, v2, text2)
    await db_session.commit()
    second = await bundle(v2)
    assert second[0]["content"] == text2 and initial[0]["content"] == text1
    assert not {n["node_id"] for n in initial[-2]["nodes"]} & {
        n["node_id"] for n in second[-2]["nodes"]
    }
    assert await bundle(v1) == initial
    assert await bundle(v2) == second
    assert job1.document_version_id == v1.id and job2.document_version_id == v2.id
    monkeypatch.setattr(
        "app.modules.processing.service.PARSER_VERSION", "future-parser-v99"
    )
    monkeypatch.setattr(
        "app.modules.processing.service.extract_document",
        lambda *_: pytest.fail("Historical version was reparsed"),
    )
    historical = await DocumentProcessingService(db_session).process_document(
        document, text1.encode(), str(v1.id)
    )
    assert (
        historical.id == p1.id
        and historical.structure_fingerprint == p1.structure_fingerprint
    )
    assert historical.parser_version == "document-structure-v1"
    assert p2.id != p1.id
    with pytest.raises(ValueError, match="fingerprint"):
        await DocumentProcessingService(db_session).process_document(
            document, text2.encode(), str(v1.id)
        )


async def test_target_and_job_reject_tenant_version_and_payload_mismatches(
    db_session, client
):
    org, user = await identity(db_session)
    foreign_org, foreign_user = await identity(db_session)
    document, version = await version_input(db_session, org, user, "Own isolated text.")
    foreign, foreign_version = await version_input(
        db_session, foreign_org, foreign_user, "Foreign isolated text."
    )
    app.dependency_overrides[get_current_active_user] = lambda: user
    await db_session.commit()
    for path in ("content", "structure"):
        assert (
            await client.get(
                f"/api/v1/documents/{document.id}/{path}",
                params={"document_version_id": foreign_version.id},
            )
        ).status_code == 404
        assert (
            await client.get(f"/api/v1/documents/{foreign.id}/{path}")
        ).status_code == 404
    with pytest.raises(HTTPException) as failure:
        await JobService(db_session).create_job(
            JobCreate(
                job_type=JobType.AI_DETECTION,
                document_id=str(document.id),
                document_version_id=str(foreign_version.id),
            ),
            str(org.id),
            str(user.id),
        )
    assert getattr(failure.value, "status_code", None) == 404
    job = Job(
        organization_id=str(org.id),
        document_id=str(document.id),
        document_version_id=str(version.id),
        job_type=JobType.DOCUMENT_PROCESSING,
        input_data={
            "document_version_id": str(version.id),
            "storage_key": version.storage_path,
        },
    )
    validate_job_target(
        job, str(org.id), str(document.id), str(version.id), version.storage_path
    )
    for triple in (
        (str(org.id), str(document.id), None),
        (str(foreign_org.id), str(document.id), str(version.id)),
        (str(org.id), str(document.id), str(foreign_version.id)),
    ):
        with pytest.raises(ValueError):
            validate_job_target(job, *triple, version.storage_path)


async def test_evidence_edges_cannot_connect_two_versions_of_one_document(db_session):
    org, user = await identity(db_session)
    document, v1 = await version_input(db_session, org, user, "Version one.")
    _, v2 = await version_input(db_session, org, user, "Version two.", document)
    service = EvidenceService(db_session)
    nodes = []
    for version in (v1, v2):
        nodes.append(
            await service.add_node(
                str(document.id),
                EvidenceNodeType.DOCUMENT,
                "Version node",
                description="Scoped identity",
                document_version_id=str(version.id),
                entity_id=str(document.id),
            )
        )
    with pytest.raises(ValueError, match="cannot cross analysis document versions"):
        await service.add_edge(
            str(nodes[0].node_id), str(nodes[1].node_id), EvidenceEdgeType.DERIVED_FROM
        )
    with pytest.raises(ValueError, match="explicit immutable"):
        await service.add_node(
            str(document.id),
            EvidenceNodeType.CLAIM,
            "Unscoped",
            description="Missing version",
            entity_id="unscoped",
        )


async def test_writing_history_defaults_to_latest_version_without_old_edits(
    db_session, client
):
    from app.modules.aegiswrite.service import AegisWriteService
    from app.modules.aegiswrite.models import EditType

    org, user = await identity(db_session)
    app.dependency_overrides[get_current_active_user] = lambda: user
    text = "In order to improve clarity, the the sentence is concise."
    document, v1 = await version_input(db_session, org, user, text)
    run = await AegisWriteService(db_session).refine(
        document_id=str(document.id),
        organization_id=str(org.id),
        user_id=str(user.id),
        text=text,
        document_version_id=str(v1.id),
        edit_types=[EditType.GRAMMAR],
        preserve_voice=True,
    )
    assert run.edits
    _, v2 = await version_input(
        db_session, org, user, "Independent second draft.", document
    )
    await db_session.commit()
    latest = (await client.get(f"/api/v1/aegiswrite/history/{document.id}")).json()
    assert latest["document_version_id"] == v2.id and latest["items"] == []
    old = (
        await client.get(
            f"/api/v1/aegiswrite/history/{document.id}",
            params={"document_version_id": v1.id},
        )
    ).json()
    assert len(old["items"]) == len(run.edits)
    assert all(item["document_version_id"] == v1.id for item in old["items"])


async def test_detection_segment_order_survives_different_database_scan_order(
    db_session,
):
    from sqlalchemy import text as sql_text
    from app.modules.detection.models import DetectionResult

    org, user = await identity(db_session)
    content = "A first sentence explains the method. A second sentence describes the sample. A third sentence records the result."
    document, version = await version_input(db_session, org, user, content)
    await analyze(db_session, org, user, document, version, content)
    document_id = str(document.id)
    for setting in ("ON", "OFF"):
        await db_session.execute(
            sql_text("PRAGMA reverse_unordered_selects=" + setting)
        )
        db_session.expire_all()
        result = await db_session.scalar(
            select(DetectionResult).where(DetectionResult.document_id == document_id)
        )
        indices = [segment.segment_index for segment in result.segments]
        assert len(indices) > 1 and indices == sorted(indices)

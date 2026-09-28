"""Exercise the final production filesystem, including native document libraries.

Run with docker run --rm -i IMAGE python - < backend/scripts/runtime_smoke.py.
Does not require test dependencies, a shell, credentials, or network access.
"""

import importlib.util
import io
import json
import os
from pathlib import Path
import platform
import shutil
import ssl
from zoneinfo import ZoneInfo
import zlib

import asyncpg
import bcrypt
from cryptography.fernet import Fernet
from docx import Document
from minio import Minio
import numpy
import pdfplumber
import pypdfium2
from celery import Celery
from fastapi import FastAPI

assert platform.python_version_tuple()[:2] == ("3", "12")
assert os.getuid() != 0
for tool in (
    "sh",
    "bash",
    "apt",
    "apt-get",
    "dpkg",
    "pip",
    "gcc",
    "cc",
    "git",
    "curl",
    "wget",
    "perl",
    "mount",
):
    assert shutil.which(tool) is None, tool
assert importlib.util.find_spec("pytest") is None
assert importlib.util.find_spec("pip") is None
assert ssl.create_default_context().get_ca_certs()
assert str(ZoneInfo("Asia/Kolkata")) == "Asia/Kolkata"
assert zlib.decompress(zlib.compress(b"verity")) == b"verity"
assert bcrypt.checkpw(b"fixture", bcrypt.hashpw(b"fixture", bcrypt.gensalt()))
box = Fernet(Fernet.generate_key())
assert box.decrypt(box.encrypt(b"fixture")) == b"fixture"
doc = Document()
doc.add_paragraph("Verity document smoke test")
buffer = io.BytesIO()
doc.save(buffer)
buffer.seek(0)
assert Document(buffer).paragraphs[0].text == "Verity document smoke test"
pdf = pypdfium2.PdfDocument.new()
pdf.new_page(100, 100)
stream = io.BytesIO()
pdf.save(stream)
stream.seek(0)
with pdfplumber.open(stream) as parsed:
    assert len(parsed.pages) == 1
assert numpy.asarray([1, 2]).sum() == 3
assert FastAPI().openapi_version == "3.1.0"
assert Celery("smoke").main == "smoke"
assert Minio("localhost:9000", secure=False)
assert asyncpg.__version__
assert not Path("/usr/local/include/python3.12").exists()
print(
    json.dumps(
        {
            "status": "PASS",
            "python": platform.python_version(),
            "machine": platform.machine(),
            "uid": os.getuid(),
            "checks": [
                "no-shell-or-package-manager",
                "no-tests-or-headers",
                "CA-certificates",
                "timezone",
                "zlib",
                "bcrypt",
                "cryptography",
                "DOCX",
                "PDF",
                "numpy",
                "FastAPI",
                "Celery",
                "asyncpg",
                "MinIO",
            ],
        }
    )
)
